import os
import io
import socket
from datetime import datetime, date, timedelta
from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
from werkzeug.security import check_password_hash
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import db_service
from db_service import DB_TYPE, MONGO_URI
from importer import preview_excel_file, import_students_from_excel

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'jim_magis_cafe_secret_key_2026_academic')
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16 MB

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Initialize database on startup (MongoDB Atlas or SQLite fallback)
db_service.init_db()

def get_local_ip():
    """Get LAN IP address for mobile access."""
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('8.8.8.8', 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return '127.0.0.1'

def get_current_meal_session():
    """Detect current meal session based on local time."""
    now = datetime.now()
    minutes = now.hour * 60 + now.minute
    
    if 420 <= minutes < 630:
        return 'Breakfast'
    elif 630 <= minutes < 930:
        return 'Lunch'
    elif 930 <= minutes < 1140:
        return 'Snacks'
    elif 1140 <= minutes < 1380:
        return 'Dinner'
    else:
        return 'Dinner' if minutes >= 1140 else 'Breakfast'

# ----------------- PAGE ROUTES & AUTH GUARDS ----------------- #

@app.route('/login')
def login_page():
    if session.get('user_id'):
        return redirect(url_for('index'))
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login_page'))

@app.route('/')
def index():
    if not session.get('user_id'):
        return redirect(url_for('login_page'))
    return render_template('index.html')

@app.route('/report/print')
def print_report():
    if not session.get('user_id'):
        return redirect(url_for('login_page'))

    req_date = request.args.get('date', date.today().isoformat())
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')
    meal = request.args.get('meal_session', '')
    violation = request.args.get('violation_type', '')
    batch = request.args.get('batch', '')

    logs, summary = db_service.get_food_logs(
        date=req_date if not (start_date and end_date) else '',
        start_date=start_date,
        end_date=end_date,
        meal_session=meal,
        violation_type=violation,
        batch=batch,
        limit=1000
    )

    date_title = f"{start_date} to {end_date}" if start_date and end_date else (req_date or "All Time")

    return render_template(
        'print_report.html',
        logs=logs,
        date_title=date_title,
        meal=meal or 'All Meals',
        batch=batch or 'All Batches',
        total_logs=summary['total'],
        hand_count=summary['eat_in_hand'] + summary['both'],
        waste_count=summary['waste_food'] + summary['both'],
        unique_students=summary['unique_students'],
        generated_at=datetime.now().strftime('%d %b %Y, %I:%M %p')
    )

# ----------------- AUTHENTICATION APIS ----------------- #

@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.get_json() or {}
    username = data.get('username', '').strip()
    password = data.get('password', '')

    if not username or not password:
        return jsonify({'error': 'Please provide both username and password.'}), 400

    user = db_service.get_user_by_username(username)
    if not user or not check_password_hash(user['password_hash'], password):
        return jsonify({'error': 'Invalid username or password.'}), 401

    session['user_id'] = str(user['id'])
    session['username'] = user['username']
    session['role'] = user['role']
    session['full_name'] = user['full_name']

    return jsonify({
        'success': True,
        'user': {
            'username': user['username'],
            'role': user['role'],
            'full_name': user['full_name']
        }
    })

@app.route('/api/me')
def api_me():
    if not session.get('user_id'):
        return jsonify({'is_authenticated': False}), 401
    return jsonify({
        'is_authenticated': True,
        'user_id': session.get('user_id'),
        'username': session.get('username'),
        'role': session.get('role'),
        'full_name': session.get('full_name')
    })

@app.route('/api/users', methods=['GET'])
def api_get_users():
    if session.get('role') != 'admin':
        return jsonify({'error': 'Administrator access required.'}), 403
    users = db_service.list_users()
    return jsonify(users)

@app.route('/api/users', methods=['POST'])
def api_create_user():
    if session.get('role') != 'admin':
        return jsonify({'error': 'Administrator access required.'}), 403
    data = request.get_json() or {}
    username = data.get('username', '').strip()
    password = data.get('password', '').strip()
    full_name = data.get('full_name', '').strip()
    role = data.get('role', 'food_committee').strip()

    if not username or not password or not full_name:
        return jsonify({'error': 'All fields are required.'}), 400

    try:
        user_id = db_service.create_user(username, password, full_name, role)
        return jsonify({'success': True, 'id': str(user_id), 'message': f'User {username} created successfully.'})
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/users/<user_id>', methods=['DELETE'])
def api_delete_user(user_id):
    if session.get('role') != 'admin':
        return jsonify({'error': 'Administrator access required.'}), 403
    db_service.delete_user(user_id)
    return jsonify({'success': True, 'message': 'User removed.'})

# ----------------- SYSTEM & CORE APIS ----------------- #

@app.route('/api/session-info')
def api_session_info():
    ip = get_local_ip()
    port = 5000
    return jsonify({
        'current_date': date.today().isoformat(),
        'current_time': datetime.now().strftime('%I:%M %p'),
        'detected_meal': get_current_meal_session(),
        'local_ip': ip,
        'local_port': port,
        'mobile_url': f"http://{ip}:{port}",
        'database_type': 'MongoDB Atlas' if ('mongodb+srv' in MONGO_URI) else ('MongoDB Local' if DB_TYPE == 'mongo' else 'SQLite'),
        'current_user': {
            'username': session.get('username', 'Guest'),
            'role': session.get('role', 'guest'),
            'full_name': session.get('full_name', 'Visitor')
        }
    })

@app.route('/api/batches')
def api_batches():
    batches = db_service.get_batches()
    return jsonify(batches)

@app.route('/api/students')
def api_students():
    search = request.args.get('search', '').strip()
    batch = request.args.get('batch', '').strip()
    log_date = request.args.get('date', date.today().isoformat())
    meal_session = request.args.get('meal_session', '')

    students = db_service.get_students(
        search=search,
        batch=batch,
        log_date=log_date,
        meal_session=meal_session
    )
    return jsonify(students)

@app.route('/api/students/<student_id_or_roll>/summary')
def api_student_summary(student_id_or_roll):
    summary = db_service.get_student_summary(student_id_or_roll)
    if not summary:
        return jsonify({'error': 'Student not found.'}), 404
    return jsonify(summary)

@app.route('/api/students', methods=['POST'])
def api_create_student():
    data = request.get_json() or {}
    roll_no = data.get('roll_no', '').strip().upper()
    name = data.get('name', '').strip().upper()
    batch = data.get('batch', 'I MBA').strip()
    mobile = data.get('mobile', '').strip()
    gender = data.get('gender', 'M').strip().upper()

    if not roll_no or not name:
        return jsonify({'error': 'Roll Number and Name are required'}), 400

    try:
        db_service.create_or_update_student(roll_no, name, batch, mobile, gender)
        return jsonify({'success': True, 'message': f'Student {name} ({roll_no}) saved successfully.'})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/students/<student_id>', methods=['DELETE'])
def api_delete_student(student_id):
    if session.get('role') not in ('admin', 'staff'):
        return jsonify({'error': 'Permission denied.'}), 403
    db_service.delete_student(student_id)
    return jsonify({'success': True, 'message': 'Student removed successfully.'})

@app.route('/api/logs', methods=['GET'])
def api_get_logs():
    req_date = request.args.get('date', '').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    meal = request.args.get('meal_session', '').strip()
    violation = request.args.get('violation_type', '').strip()
    batch = request.args.get('batch', '').strip()
    search = request.args.get('search', '').strip()
    limit = request.args.get('limit', 300, type=int)

    logs, summary = db_service.get_food_logs(
        date=req_date,
        start_date=start_date,
        end_date=end_date,
        meal_session=meal,
        violation_type=violation,
        batch=batch,
        search=search,
        limit=limit
    )

    return jsonify({
        'logs': logs,
        'summary': summary
    })

@app.route('/api/logs', methods=['POST'])
def api_create_logs():
    data = request.get_json() or {}
    student_ids = data.get('student_ids', [])
    log_date = data.get('log_date', date.today().isoformat())
    meal_session = data.get('meal_session', get_current_meal_session())
    violation_type = data.get('violation_type', 'waste_food')
    waste_severity = data.get('waste_severity', 'Normal')
    remarks = data.get('remarks', '').strip()
    logged_by = session.get('full_name') or 'Food Committee'

    if not student_ids:
        return jsonify({'error': 'No students selected'}), 400

    created_ids, recorded_names = db_service.create_food_logs(
        student_ids=student_ids,
        log_date=log_date,
        meal_session=meal_session,
        violation_type=violation_type,
        waste_severity=waste_severity,
        remarks=remarks,
        logged_by=logged_by
    )

    return jsonify({
        'success': True,
        'count': len(created_ids),
        'created_ids': created_ids,
        'recorded_students': recorded_names,
        'message': f"Logged {len(created_ids)} student(s) for {meal_session} on {log_date}."
    })

@app.route('/api/logs/<log_id>', methods=['DELETE'])
def api_delete_log(log_id):
    db_service.delete_food_log(log_id)
    return jsonify({'success': True, 'message': 'Log entry removed successfully.'})

@app.route('/api/analytics')
def api_analytics():
    analytics = db_service.get_analytics_data()
    return jsonify(analytics)

@app.route('/api/upload-excel', methods=['POST'])
def api_upload_excel():
    if session.get('role') not in ('admin', 'staff'):
        return jsonify({'error': 'Administrator or Staff permission required.'}), 403

    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400

    batch_name = request.form.get('batch_name', 'II MBA').strip()
    action = request.form.get('action', 'import')

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{filename}")
    file.save(filepath)

    try:
        if action == 'preview':
            preview = preview_excel_file(filepath)
            return jsonify({'success': True, 'preview': preview})
        else:
            result = import_students_from_excel(filepath, default_batch=batch_name)
            return jsonify({
                'success': True,
                'batch_name': batch_name,
                'imported': result['imported'],
                'updated': result['updated'],
                'skipped': result['skipped'],
                'errors': result['errors']
            })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/export/excel')
def export_excel():
    req_date = request.args.get('date', '')
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')
    meal = request.args.get('meal_session', '')
    violation = request.args.get('violation_type', '')
    batch = request.args.get('batch', '')
    search = request.args.get('search', '')

    logs, _ = db_service.get_food_logs(
        date=req_date,
        start_date=start_date,
        end_date=end_date,
        meal_session=meal,
        violation_type=violation,
        batch=batch,
        search=search,
        limit=2000
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Magis Cafe Food Logs"

    # Title Block
    ws.merge_cells('A1:I1')
    ws['A1'] = "JIM - MAGIS CAFE FOOD COMMITTEE MONITORING RECORD"
    ws['A1'].font = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

    ws.merge_cells('A2:I2')
    sub_title = f"Generated on: {datetime.now().strftime('%d %B %Y, %I:%M %p')} | Total Records: {len(logs)}"
    ws['A2'] = sub_title
    ws['A2'].font = Font(name='Segoe UI', size=10, italic=True, color='475569')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        "S.No", "Date", "Meal Session", "Roll Number", 
        "Student Name", "Batch", "Violation Type", "Waste Severity", "Remarks"
    ]
    header_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border_thin
    ws.row_dimensions[4].height = 26

    # Data Rows
    row_alt_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    violation_labels = {
        'eat_in_hand': 'Eat in Hand',
        'waste_food': 'Wasted Food',
        'both': 'Both (Hand + Waste)'
    }

    for idx, l in enumerate(logs, start=1):
        current_row = idx + 4
        v_label = violation_labels.get(l.get('violation_type'), l.get('violation_type', ''))
        
        row_values = [
            idx,
            l.get('log_date', ''),
            l.get('meal_session', ''),
            l.get('roll_no', ''),
            l.get('student_name', ''),
            l.get('batch', ''),
            v_label,
            l.get('waste_severity', 'Normal'),
            l.get('remarks', '-') or '-'
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = Font(name='Segoe UI', size=10)
            cell.border = border_thin
            if idx % 2 == 0:
                cell.fill = row_alt_fill
            
            if col_idx in [1, 2, 3, 4, 6, 8]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

        ws.row_dimensions[current_row].height = 22

    col_widths = [8, 14, 15, 16, 30, 12, 22, 16, 32]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Magis_Cafe_Logs_{date.today().strftime('%Y%m%d')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

if __name__ == '__main__':
    local_ip = get_local_ip()
    print("=" * 60)
    print("  JIM - MAGIS CAFE FOOD COMMITTEE MONITORING SYSTEM")
    print(f"  Database Engine: {db_service.DB_TYPE.upper()}")
    print("=" * 60)
    print(f" * Local PC Access:    http://127.0.0.1:5000")
    print(f" * Mobile Phone Access: http://{local_ip}:5000")
    print("=" * 60)
    app.run(host='0.0.0.0', port=5000, debug=True)
