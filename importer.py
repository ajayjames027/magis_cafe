import os
import openpyxl
import sqlite3
import re
from database import get_db_connection

def normalize_header(header):
    if not header:
        return ''
    return re.sub(r'[^a-z0-9]', '', str(header).strip().lower())

def detect_columns(headers):
    """
    Detect column indices for roll_no, name, mobile, gender, batch
    """
    col_map = {
        'roll_no': None,
        'name': None,
        'mobile': None,
        'gender': None,
        'batch': None
    }

    roll_aliases = ['dno', 'd_no', 'rollno', 'roll_no', 'roll', 'regno', 'reg_no', 'register', 'registerno', 'studentid', 'idno']
    name_aliases = ['candidatename', 'studentname', 'name', 'fullname', 'nameofcandidate', 'nameofstudent']
    mobile_aliases = ['mobileno', 'mobile', 'phone', 'phoneno', 'contact', 'contactno']
    gender_aliases = ['gender', 'sex']
    batch_aliases = ['batch', 'class', 'year', 'course', 'program', 'programme']

    for idx, raw_h in enumerate(headers):
        if not raw_h:
            continue
        norm_h = normalize_header(raw_h)

        if col_map['roll_no'] is None and any(alias in norm_h for alias in roll_aliases):
            col_map['roll_no'] = idx
        elif col_map['name'] is None and any(alias in norm_h for alias in name_aliases):
            col_map['name'] = idx
        elif col_map['mobile'] is None and any(alias in norm_h for alias in mobile_aliases):
            col_map['mobile'] = idx
        elif col_map['gender'] is None and any(alias in norm_h for alias in gender_aliases):
            col_map['gender'] = idx
        elif col_map['batch'] is None and any(alias in norm_h for alias in batch_aliases):
            col_map['batch'] = idx

    return col_map

def preview_excel_file(filepath):
    """
    Returns headers, detected columns, sample rows, and total estimated count.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {'error': 'The file is empty.'}

    # Find the header row (sometimes row 1 is title, row 2 is header)
    header_idx = 0
    headers = []
    for i, r in enumerate(rows[:5]):
        str_row = [str(c or '').strip() for c in r]
        # Check if row looks like a header
        if any(any(k in normalize_header(c) for k in ['dno', 'name', 'roll', 'candidate', 'sno']) for c in str_row):
            header_idx = i
            headers = str_row
            break

    if not headers:
        header_idx = 0
        headers = [str(c or '').strip() for c in rows[0]]

    col_map = detect_columns(headers)

    sample_rows = []
    data_rows = rows[header_idx + 1:]
    for r in data_rows[:5]:
        if any(r):
            sample_rows.append([str(c or '').strip() for c in r])

    valid_count = sum(1 for r in data_rows if any(r))

    return {
        'headers': headers,
        'detected_mapping': col_map,
        'sample_rows': sample_rows,
        'total_rows': valid_count,
        'sheet_names': wb.sheetnames
    }

def import_students_from_excel(filepath, default_batch='I MBA', custom_mapping=None):
    """
    Imports students from an Excel file into the SQLite database.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {'imported': 0, 'updated': 0, 'skipped': 0, 'errors': ['Empty file']}

    header_idx = 0
    headers = []
    for i, r in enumerate(rows[:5]):
        str_row = [str(c or '').strip() for c in r]
        if any(any(k in normalize_header(c) for k in ['dno', 'name', 'roll', 'candidate', 'sno']) for c in str_row):
            header_idx = i
            headers = str_row
            break

    if not headers:
        header_idx = 0
        headers = [str(c or '').strip() for c in rows[0]]

    col_map = custom_mapping if custom_mapping else detect_columns(headers)

    # Fallbacks if detection missed
    if col_map.get('roll_no') is None:
        # Default to col 1 if col 0 is SNO
        if len(headers) > 1 and normalize_header(headers[0]) in ['sno', 'slno', 'no']:
            col_map['roll_no'] = 1
        else:
            col_map['roll_no'] = 0

    if col_map.get('name') is None:
        if len(headers) > 2 and col_map.get('roll_no') == 1:
            col_map['name'] = 2
        else:
            col_map['name'] = 1

    conn = get_db_connection()
    cursor = conn.cursor()

    imported = 0
    updated = 0
    skipped = 0
    errors = []

    data_rows = rows[header_idx + 1:]

    for line_num, r in enumerate(data_rows, start=header_idx + 2):
        if not any(r):
            continue

        try:
            roll_idx = col_map.get('roll_no')
            name_idx = col_map.get('name')
            mobile_idx = col_map.get('mobile')
            gender_idx = col_map.get('gender')
            batch_idx = col_map.get('batch')

            roll_val = str(r[roll_idx] or '').strip() if roll_idx is not None and roll_idx < len(r) else ''
            name_val = str(r[name_idx] or '').strip() if name_idx is not None and name_idx < len(r) else ''

            if not roll_val or not name_val:
                skipped += 1
                continue

            mobile_val = str(r[mobile_idx] or '').strip() if mobile_idx is not None and mobile_idx < len(r) else ''
            # Handle float representation of mobile numbers from excel
            if mobile_val.endswith('.0'):
                mobile_val = mobile_val[:-2]

            gender_val = str(r[gender_idx] or '').strip().upper() if gender_idx is not None and gender_idx < len(r) else ''
            if gender_val:
                gender_val = gender_val[0]  # 'M' or 'F'

            batch_val = default_batch
            if batch_idx is not None and batch_idx < len(r) and r[batch_idx]:
                batch_val = str(r[batch_idx]).strip()

            # Insert or Update
            cursor.execute('SELECT id FROM students WHERE roll_no = ?', (roll_val,))
            existing = cursor.fetchone()

            if existing:
                cursor.execute('''
                    UPDATE students 
                    SET name = ?, batch = ?, mobile = ?, gender = ?
                    WHERE id = ?
                ''', (name_val, batch_val, mobile_val, gender_val, existing['id']))
                updated += 1
            else:
                cursor.execute('''
                    INSERT INTO students (roll_no, name, batch, mobile, gender)
                    VALUES (?, ?, ?, ?, ?)
                ''', (roll_val, name_val, batch_val, mobile_val, gender_val))
                imported += 1

        except Exception as e:
            errors.append(f"Row {line_num}: {str(e)}")

    conn.commit()
    conn.close()

    return {
        'imported': imported,
        'updated': updated,
        'skipped': skipped,
        'errors': errors
    }
